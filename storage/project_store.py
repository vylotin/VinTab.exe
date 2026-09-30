import sqlite3
import json
from contextlib import closing
from datetime import date, datetime
import math
from pathlib import Path


TABELA_METADADOS = "__vintab_metadata"
TABELA_FORMATOS = "__vintab_datasets"
TABELA_GRAFICOS = "__vintab_charts"


def listar_conjuntos(caminho_banco):
    caminho = Path(caminho_banco).resolve()
    uri = caminho.as_uri() + "?mode=ro"
    with closing(sqlite3.connect(uri, uri=True)) as conexao:
        resultados = conexao.execute(
            "SELECT name FROM sqlite_master "
            "WHERE type = 'table' AND name NOT LIKE 'sqlite_%' "
            "ORDER BY name"
        ).fetchall()
    return [
        resultado[0] for resultado in resultados
        if not resultado[0].startswith("__vintab_")
    ]


def ler_conjunto(caminho_banco, nome_conjunto):
    caminho = Path(caminho_banco).resolve()
    uri = caminho.as_uri() + "?mode=ro"
    with closing(sqlite3.connect(uri, uri=True)) as conexao:
        colunas = _obter_colunas(conexao, nome_conjunto)
        linhas = conexao.execute(
            f"SELECT * FROM {_identificador(nome_conjunto)}"
        ).fetchall()
    return colunas, linhas


def listar_colunas_conjunto(caminho_banco, nome_conjunto):
    caminho = Path(caminho_banco).resolve()
    uri = caminho.as_uri() + "?mode=ro"
    with closing(sqlite3.connect(uri, uri=True)) as conexao:
        return _obter_colunas(conexao, nome_conjunto)


def criar_conjunto(caminho_projeto, nome_conjunto, colunas):
    caminho = _validar_projeto(caminho_projeto)
    nome = _validar_nome(nome_conjunto, "nome do conjunto")
    colunas = _validar_colunas(colunas)
    definicoes = ", ".join(f"{_identificador(coluna)} NUMERIC" for coluna in colunas)

    with closing(sqlite3.connect(caminho)) as conexao:
        with conexao:
            conexao.execute(
                f"CREATE TABLE {_identificador(nome)} ({definicoes})"
            )


def criar_projeto(caminho_projeto, formato, nome_conjunto, colunas):
    caminho = Path(caminho_projeto).expanduser().resolve()
    if caminho.suffix.lower() != ".vt":
        raise ValueError("O projeto deve usar a extensão .vt.")
    if caminho.exists():
        raise FileExistsError(f"Já existe um arquivo neste local: {caminho}")
    if not caminho.parent.is_dir():
        raise FileNotFoundError(f"A pasta do projeto não existe: {caminho.parent}")

    formato = _validar_formato(formato)
    nome_conjunto = _validar_nome(nome_conjunto, "nome do conjunto")
    colunas = _validar_colunas(colunas)
    try:
        with closing(sqlite3.connect(caminho)) as conexao:
            with conexao:
                _criar_tabelas_internas(conexao)
                conexao.execute(
                    f"INSERT INTO {_identificador(TABELA_METADADOS)} (chave, valor) "
                    "VALUES ('format', ?)",
                    (formato,),
                )
                definicoes = ", ".join(
                    f"{_identificador(coluna)} NUMERIC" for coluna in colunas
                )
                conexao.execute(
                    f"CREATE TABLE {_identificador(nome_conjunto)} ({definicoes})"
                )
                conexao.execute(
                    f"INSERT INTO {_identificador(TABELA_FORMATOS)} (nome, formato) "
                    "VALUES (?, ?)",
                    (nome_conjunto, formato),
                )
    except Exception:
        caminho.unlink(missing_ok=True)
        raise
    return caminho


def obter_formato_projeto(caminho_projeto):
    caminho = Path(caminho_projeto).resolve()
    with closing(sqlite3.connect(caminho.as_uri() + "?mode=ro", uri=True)) as conexao:
        tabelas = {
            nome for (nome,) in conexao.execute(
                "SELECT name FROM sqlite_master WHERE type = 'table'"
            )
        }
        if TABELA_METADADOS in tabelas:
            registro = conexao.execute(
                f"SELECT valor FROM {_identificador(TABELA_METADADOS)} "
                "WHERE chave = 'format'"
            ).fetchone()
            if registro:
                return _validar_formato(registro[0])

        nomes_dados = [
            nome for nome in tabelas
            if not nome.startswith("sqlite_") and not nome.startswith("__vintab_")
        ]
        for nome in nomes_dados:
            colunas = _obter_colunas(conexao, nome)
            if any(
                coluna in {"Mês/ano", "ISC-NEURO", "ISC-Coluna", "ISC-Quadril"}
                for coluna in colunas
            ):
                return "ISC"
    return "IRAS"


def obter_formato_conjunto(caminho_projeto, nome_conjunto):
    caminho = Path(caminho_projeto).resolve()
    with closing(sqlite3.connect(caminho.as_uri() + "?mode=ro", uri=True)) as conexao:
        tabelas = {
            nome for (nome,) in conexao.execute(
                "SELECT name FROM sqlite_master WHERE type = 'table'"
            )
        }
        if TABELA_FORMATOS in tabelas:
            registro = conexao.execute(
                f"SELECT formato FROM {_identificador(TABELA_FORMATOS)} WHERE nome = ?",
                (nome_conjunto,),
            ).fetchone()
            if registro:
                return _validar_formato(registro[0])
    return obter_formato_projeto(caminho)


def registrar_formato_conjunto(caminho_projeto, nome_conjunto, formato):
    caminho = _validar_projeto(caminho_projeto)
    nome_conjunto = _validar_nome(nome_conjunto, "nome do conjunto")
    formato = _validar_formato(formato)
    with closing(sqlite3.connect(caminho)) as conexao:
        with conexao:
            _criar_tabelas_internas(conexao)
            conexao.execute(
                f"INSERT INTO {_identificador(TABELA_FORMATOS)} (nome, formato) "
                "VALUES (?, ?) ON CONFLICT(nome) DO UPDATE SET formato = excluded.formato",
                (nome_conjunto, formato),
            )


def listar_graficos(caminho_projeto):
    caminho = Path(caminho_projeto).resolve()
    with closing(sqlite3.connect(caminho.as_uri() + "?mode=ro", uri=True)) as conexao:
        existe = conexao.execute(
            "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = ?",
            (TABELA_GRAFICOS,),
        ).fetchone()
        if not existe:
            return []
        linhas = conexao.execute(
            f"SELECT id, nome, conjunto, configuracao, criado_em "
            f"FROM {_identificador(TABELA_GRAFICOS)} ORDER BY criado_em, id"
        ).fetchall()
    return [
        {
            "id": identificador,
            "name": nome,
            "dataset": conjunto,
            "config": json.loads(configuracao),
            "created_at": criado_em,
        }
        for identificador, nome, conjunto, configuracao, criado_em in linhas
    ]


def salvar_grafico(caminho_projeto, nome, conjunto, configuracao):
    caminho = _validar_projeto(caminho_projeto)
    nome = _validar_nome(nome, "nome do gráfico")
    conjunto = _validar_nome(conjunto, "nome do conjunto")
    if conjunto not in listar_conjuntos(caminho):
        raise ValueError(f"Conjunto não encontrado no projeto: {conjunto}")
    conteudo = json.dumps(configuracao, ensure_ascii=False, allow_nan=False)
    with closing(sqlite3.connect(caminho)) as conexao:
        with conexao:
            _criar_tabelas_internas(conexao)
            cursor = conexao.execute(
                f"INSERT INTO {_identificador(TABELA_GRAFICOS)} "
                "(nome, conjunto, configuracao) VALUES (?, ?, ?)",
                (nome, conjunto, conteudo),
            )
            return cursor.lastrowid


def excluir_grafico(caminho_projeto, identificador):
    caminho = _validar_projeto(caminho_projeto)
    with closing(sqlite3.connect(caminho)) as conexao:
        with conexao:
            cursor = conexao.execute(
                f"DELETE FROM {_identificador(TABELA_GRAFICOS)} WHERE id = ?",
                (int(identificador),),
            )
            return cursor.rowcount > 0


def salvar_conjunto(caminho_projeto, nome_conjunto, colunas, linhas):
    caminho = _validar_projeto(caminho_projeto)
    nome = _validar_nome(nome_conjunto, "nome do conjunto")
    colunas = _validar_colunas(colunas)
    tabela_sql = _identificador(nome)
    colunas_existentes = _ler_colunas_projeto(caminho, nome)
    if colunas_existentes != colunas:
        raise ValueError("As colunas do conjunto foram alteradas durante a edição.")

    linhas_validas = []
    for linha in linhas:
        if len(linha) != len(colunas):
            raise ValueError("Uma linha contém quantidade incorreta de valores.")
        valores = [
            None if valor is None or not str(valor).strip() else str(valor)
            for valor in linha
        ]
        if any(valor not in (None, "") for valor in valores):
            linhas_validas.append(valores)

    marcadores = ", ".join("?" for _ in colunas)
    nomes_colunas = ", ".join(_identificador(coluna) for coluna in colunas)
    with closing(sqlite3.connect(caminho)) as conexao:
        with conexao:
            conexao.execute(f"DELETE FROM {tabela_sql}")
            if linhas_validas:
                conexao.executemany(
                    f"INSERT INTO {tabela_sql} ({nomes_colunas}) VALUES ({marcadores})",
                    linhas_validas,
                )


def importar_planilha_excel(caminho_planilha, caminho_projeto):
    origem = Path(caminho_planilha).resolve()
    destino = Path(caminho_projeto).resolve()
    if origem.suffix.lower() not in {".xlsx", ".xls"} or not origem.is_file():
        raise ValueError("Selecione uma planilha Excel .xlsx ou .xls existente.")
    if destino.suffix.lower() != ".vt":
        raise ValueError("O destino da importação precisa ser um projeto .vt.")

    novo_projeto = not destino.exists()
    if not novo_projeto and not destino.is_file():
        raise ValueError("O destino selecionado não é um arquivo válido.")

    try:
        planilhas = _ler_planilhas_excel(origem)
        if not planilhas:
            raise ValueError("A planilha não contém abas para importar.")

        resumo = []
        with closing(sqlite3.connect(destino)) as conexao:
            with conexao:
                tabelas_existentes = {
                    nome.casefold()
                    for (nome,) in conexao.execute(
                        "SELECT name FROM sqlite_master WHERE type = 'table'"
                    )
                }
                for nome_aba, (colunas_originais, linhas_originais) in planilhas.items():
                    nome_tabela = _validar_nome(nome_aba, "nome da aba")
                    if nome_tabela.casefold() in tabelas_existentes:
                        raise ValueError(
                            f"A tabela '{nome_tabela}' já existe no projeto. "
                            "Renomeie a aba ou escolha outro projeto."
                        )

                    colunas = _validar_colunas(colunas_originais)
                    tabela_sql = _identificador(nome_tabela)
                    definicoes = ", ".join(
                        f"{_identificador(coluna)} NUMERIC" for coluna in colunas
                    )
                    conexao.execute(
                        f"CREATE TABLE {tabela_sql} ({definicoes})"
                    )

                    marcadores = ", ".join("?" for _ in colunas)
                    nomes_colunas = ", ".join(
                        _identificador(coluna) for coluna in colunas
                    )
                    linhas = linhas_originais
                    if linhas:
                        conexao.executemany(
                            f"INSERT INTO {tabela_sql} ({nomes_colunas}) "
                            f"VALUES ({marcadores})",
                            linhas,
                        )
                    tabelas_existentes.add(nome_tabela.casefold())
                    resumo.append((nome_tabela, len(linhas)))
        return resumo
    except Exception:
        if novo_projeto and destino.exists():
            destino.unlink()
        raise


def _ler_planilhas_excel(caminho):
    planilhas = {}
    if caminho.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        livro = load_workbook(caminho, read_only=True, data_only=True)
        try:
            for planilha in livro.worksheets:
                linhas = planilha.iter_rows(values_only=True)
                cabecalho = next(linhas, None)
                if cabecalho is None:
                    raise ValueError(f"A aba '{planilha.title}' está vazia.")
                cabecalho = list(cabecalho)
                while cabecalho and cabecalho[-1] is None:
                    cabecalho.pop()
                if not cabecalho or any(
                    valor is None or not str(valor).strip() for valor in cabecalho
                ):
                    raise ValueError(
                        f"A aba '{planilha.title}' precisa ter cabeçalhos em todas as colunas."
                    )
                colunas = [str(valor).strip() for valor in cabecalho]
                linhas_convertidas = []
                for linha in linhas:
                    valores = list(linha)
                    if any(valor is not None for valor in valores[len(colunas):]):
                        raise ValueError(
                            f"A aba '{planilha.title}' contém dados sem cabeçalho."
                        )
                    valores = valores[:len(colunas)]
                    valores.extend([None] * (len(colunas) - len(valores)))
                    linhas_convertidas.append(
                        [_valor_excel(valor) for valor in valores]
                    )
                planilhas[planilha.title] = colunas, linhas_convertidas
        finally:
            livro.close()
        return planilhas

    import xlrd

    livro = xlrd.open_workbook(str(caminho), on_demand=True)
    try:
        for planilha in livro.sheets():
            if planilha.nrows == 0:
                raise ValueError(f"A aba '{planilha.name}' está vazia.")
            cabecalho = planilha.row_values(0)
            while cabecalho and cabecalho[-1] == "":
                cabecalho.pop()
            if not cabecalho or any(not str(valor).strip() for valor in cabecalho):
                raise ValueError(
                    f"A aba '{planilha.name}' precisa ter cabeçalhos em todas as colunas."
                )
            colunas = [str(valor).strip() for valor in cabecalho]
            linhas_convertidas = []
            for indice_linha in range(1, planilha.nrows):
                linha = planilha.row(indice_linha)
                if any(cell.ctype != xlrd.XL_CELL_EMPTY for cell in linha[len(colunas):]):
                    raise ValueError(
                        f"A aba '{planilha.name}' contém dados sem cabeçalho."
                    )
                valores = []
                for cell in linha[:len(colunas)]:
                    valor = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        valor = xlrd.xldate.xldate_as_datetime(valor, livro.datemode)
                    elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        valor = None
                    valores.append(_valor_excel(valor))
                valores.extend([None] * (len(colunas) - len(valores)))
                linhas_convertidas.append(valores)
            planilhas[planilha.name] = colunas, linhas_convertidas
    finally:
        livro.release_resources()
    return planilhas


def _valor_excel(valor):
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, datetime):
        return valor.isoformat(sep=" ")
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, float) and math.isnan(valor):
        return None
    if hasattr(valor, "item"):
        valor = valor.item()
    if isinstance(valor, (str, int, float, bytes)):
        return valor
    return str(valor)


def _ler_colunas_projeto(caminho_projeto, nome_conjunto):
    with closing(sqlite3.connect(caminho_projeto)) as conexao:
        return _obter_colunas(conexao, nome_conjunto)


def _validar_formato(formato):
    valor = str(formato).strip().upper()
    if valor not in {"IRAS", "ISC"}:
        raise ValueError("O formato do projeto deve ser IRAS ou ISC.")
    return valor


def _criar_tabelas_internas(conexao):
    conexao.execute(
        f"CREATE TABLE IF NOT EXISTS {_identificador(TABELA_METADADOS)} ("
        "chave TEXT PRIMARY KEY, valor TEXT NOT NULL)"
    )
    conexao.execute(
        f"CREATE TABLE IF NOT EXISTS {_identificador(TABELA_FORMATOS)} ("
        "nome TEXT PRIMARY KEY, formato TEXT NOT NULL)"
    )
    conexao.execute(
        f"CREATE TABLE IF NOT EXISTS {_identificador(TABELA_GRAFICOS)} ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT, "
        "nome TEXT NOT NULL, conjunto TEXT NOT NULL, "
        "configuracao TEXT NOT NULL, criado_em TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP)"
    )


def _obter_colunas(conexao, nome_conjunto):
    nome = _validar_nome(nome_conjunto, "nome do conjunto")
    resultado = conexao.execute(
        f"PRAGMA table_info({_identificador(nome)})"
    ).fetchall()
    if not resultado:
        raise ValueError(f"Conjunto não encontrado: {nome}")
    return [coluna[1] for coluna in resultado]


def _validar_projeto(caminho_projeto):
    caminho = Path(caminho_projeto).resolve()
    if caminho.suffix.lower() != ".vt":
        raise ValueError("A edição de dados está disponível somente em projetos .vt.")
    if not caminho.is_file():
        raise FileNotFoundError(f"Projeto não encontrado: {caminho}")
    return caminho


def _validar_colunas(colunas):
    nomes = [_validar_nome(coluna, "nome da coluna") for coluna in colunas]
    if not nomes:
        raise ValueError("Informe pelo menos uma coluna.")
    if len({nome.casefold() for nome in nomes}) != len(nomes):
        raise ValueError("Os nomes das colunas não podem se repetir.")
    return nomes


def _validar_nome(valor, descricao):
    nome = str(valor).strip()
    if not nome or "\x00" in nome:
        raise ValueError(f"Informe um {descricao} válido.")
    if nome.casefold().startswith("sqlite_"):
        raise ValueError("Nomes iniciados por 'sqlite_' são reservados pelo SQLite.")
    return nome


def _identificador(valor):
    return '"' + valor.replace('"', '""') + '"'